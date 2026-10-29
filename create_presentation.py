from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = "GMU_Smart_Campus_PBL_Presentation.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

NAVY = "172554"
BLUE = "2563EB"
PURPLE = "7C3AED"
CYAN = "06B6D4"
GREEN = "16A34A"
AMBER = "F59E0B"
RED = "DC2626"
WHITE = "FFFFFF"
LIGHT = "F1F5F9"
INK = "1E293B"
MUTED = "64748B"
PALE_BLUE = "DBEAFE"
PALE_CYAN = "CFFAFE"
PALE_PURPLE = "EDE9FE"
PALE_GREEN = "DCFCE7"
PALE_AMBER = "FEF3C7"
PALE_RED = "FEE2E2"
FONT = "Aptos"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def rect(slide, x, y, w, h, color, radius=False, line=None, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def circle(slide, x, y, d, color, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line:
        shape.line.color.rgb = rgb(line)
    else:
        shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def bullets(slide, items, x, y, w, h, size=17, color=INK, gap=9, bullet_color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.level = 0
        p.space_after = Pt(gap)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
    return box


def line(slide, x1, y1, x2, y2, color=BLUE, width=2):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return sh


def pill(slide, label, x, y, w, color=BLUE, pale=PALE_BLUE):
    rect(slide, x, y, w, 0.32, pale, radius=True)
    text(slide, label, x + 0.08, y + 0.025, w - 0.16, 0.25, 10, color, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def card(slide, x, y, w, h, title, body="", accent=BLUE, pale=PALE_BLUE, title_size=17, body_size=13):
    rect(slide, x, y, w, h, WHITE, radius=True)
    rect(slide, x, y, 0.08, h, accent, radius=True)
    text(slide, title, x + 0.22, y + 0.16, w - 0.4, 0.38, title_size, NAVY, True)
    if body:
        text(slide, body, x + 0.22, y + 0.61, w - 0.42, h - 0.72, body_size, INK)


def footer(slide, n, dark=False):
    c = "CBD5E1" if dark else MUTED
    text(slide, "GM UNIVERSITY  •  BCA  •  UC25CA2302", 0.42, 7.15, 5.2, 0.18, 8, c, True, margin=0)
    text(slide, f"{n:02d} / 20", 12.13, 7.12, 0.78, 0.2, 9, c, True, align=PP_ALIGN.RIGHT, margin=0)


def new_slide(title, n, kicker="SMART CAMPUS PBL"):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(LIGHT)
    rect(s, 0, 0, 13.333, 0.13, BLUE)
    text(s, kicker, 0.55, 0.32, 4.8, 0.22, 9, BLUE, True, margin=0)
    text(s, title, 0.55, 0.63, 12.1, 0.55, 25, NAVY, True, margin=0)
    rect(s, 0.55, 1.27, 0.72, 0.055, CYAN)
    footer(s, n)
    return s


def node(slide, x, y, w, h, label, color, pale, size=12):
    rect(slide, x, y, w, h, pale, radius=True)
    rect(slide, x, y, 0.07, h, color, radius=True)
    text(slide, label, x + 0.13, y + 0.04, w - 0.2, h - 0.08, size, NAVY, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def label_value(slide, label, value, x, y, w):
    text(slide, label, x, y, w, 0.22, 10, MUTED, True, margin=0)
    text(slide, value, x, y + 0.24, w, 0.34, 15, NAVY, True, margin=0)


# 1. Cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(NAVY)
rect(s, 8.5, 0, 4.833, 7.5, BLUE)
rect(s, 8.5, 0, 4.833, 7.5, PURPLE, transparency=45)
for x, y, d, c in [(9.35, 0.9, .18, CYAN), (11.8, 1.55, .15, WHITE), (10.4, 2.4, .18, AMBER),
                   (12.2, 3.2, .16, CYAN), (9.25, 4.1, .14, WHITE), (11.1, 5.25, .2, GREEN)]:
    circle(s, x, y, d, c)
for a, b, c, d in [(9.44, 1.0, 10.48, 2.5), (10.5, 2.5, 11.88, 1.62), (10.5, 2.5, 12.28, 3.3),
                   (10.5, 2.5, 9.32, 4.2), (9.32, 4.2, 11.2, 5.35), (11.2, 5.35, 12.28, 3.3)]:
    line(s, a, b, c, d, "BFDBFE", 1.2)
# Campus icon made of editable shapes
rect(s, 9.15, 3.55, 3.35, 1.1, WHITE, radius=True)
rect(s, 9.47, 3.0, 2.7, 0.56, "E0F2FE", radius=True)
for x in [9.58, 10.2, 10.82, 11.44]:
    rect(s, x, 3.75, .32, .36, BLUE, radius=True)
rect(s, 10.6, 4.05, .42, .6, NAVY, radius=True)
text(s, "GM UNIVERSITY", 0.7, 0.62, 5.4, 0.3, 12, "93C5FD", True, margin=0)
text(s, "AI-Enabled Smart Campus\nMonitoring and Intelligent\nStudent Support System", 0.7, 1.28, 7.6, 2.0, 30, WHITE, True, margin=0)
text(s, "Using Python, NLP and Data Analytics", 0.72, 3.47, 6.7, 0.45, 19, "BAE6FD", False, margin=0)
rect(s, 0.72, 4.22, 5.65, 0.04, CYAN)
text(s, "BCA  •  Semester 3  •  Course Code UC25CA2302", 0.72, 4.48, 6.7, 0.32, 13, WHITE, True, margin=0)
text(s, "ACADEMIC PBL PROJECT", 0.72, 5.0, 4.0, 0.28, 10, "C4B5FD", True, margin=0)
text(s, "Team members: [Add names]     USNs: [Add USNs]\nFaculty guide: [Add name]     Academic year: [Add year]",
     0.72, 5.42, 6.6, 0.8, 12, "E2E8F0", False, margin=0)
footer(s, 1, True)

# 2. Overview
s = new_slide("Project overview", 2)
text(s, "A Python-based concept for bringing campus information together and helping staff review student support needs.",
     0.65, 1.57, 11.9, 0.55, 17, INK, False)
src = [("Student records", BLUE, PALE_BLUE), ("Attendance", CYAN, PALE_CYAN),
       ("Marks", PURPLE, PALE_PURPLE), ("Feedback", AMBER, PALE_AMBER),
       ("Campus observations", GREEN, PALE_GREEN)]
for i, (lab, col, pale) in enumerate(src):
    y = 2.35 + i * 0.72
    node(s, .8, y, 2.5, .47, lab, col, pale, 13)
    line(s, 3.38, y + .24, 5.45, 3.95, "94A3B8", 1.4)
rect(s, 5.48, 3.1, 2.55, 1.55, NAVY, radius=True)
text(s, "SMART CAMPUS\nSYSTEM", 5.7, 3.45, 2.1, .7, 17, WHITE, True,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line(s, 8.05, 3.87, 9.0, 3.87, BLUE, 2)
node(s, 9.15, 2.55, 3.2, .65, "Analyse information", BLUE, PALE_BLUE, 14)
node(s, 9.15, 3.55, 3.2, .65, "Suggest support actions", GREEN, PALE_GREEN, 14)
node(s, 9.15, 4.55, 3.2, .65, "Create reports for review", PURPLE, PALE_PURPLE, 14)
text(s, "Decision support for staff — not an automated decision-maker.", 5.0, 5.55, 7.3, .35, 13, MUTED, True,
     align=PP_ALIGN.CENTER)

# 3. Problem
s = new_slide("Problem statement", 3)
text(s, "Project problem areas to address — not verified findings about GM University.", .65, 1.55, 11.9, .35, 13, MUTED, True)
rect(s, .72, 2.14, 5.6, 3.8, WHITE, radius=True)
rect(s, 7.0, 2.14, 5.6, 3.8, NAVY, radius=True)
pill(s, "SEPARATE / MANUAL VIEW", .98, 2.4, 2.45, AMBER, PALE_AMBER)
text(s, "Records may sit in separate places", 1.0, 2.92, 4.9, .45, 18, NAVY, True)
bullets(s, ["Attendance and marks need regular monitoring.",
            "Feedback and concerns need systematic review.",
            "Repeated concerns can be difficult to spot manually.",
            "Staff need timely, useful information."], 1.02, 3.52, 4.9, 2.0, 15, INK, 10)
pill(s, "PROPOSED INTEGRATED VIEW", 7.28, 2.4, 2.5, CYAN, PALE_CYAN)
text(s, "One organised review workflow", 7.3, 2.92, 4.9, .45, 18, WHITE, True)
for i, lab in enumerate(["Collect", "Organise", "Analyse", "Review & support"]):
    x = 7.35 + (i % 2) * 2.4
    y = 3.72 + (i // 2) * 1.0
    node(s, x, y, 2.0, .56, lab, CYAN if i % 2 == 0 else PURPLE, "E0F2FE", 13)
text(s, "Better visibility while keeping staff in the loop.", 7.35, 5.65, 4.8, .26, 12, "CBD5E1", True)

# 4. Objectives
s = new_slide("Project objectives", 4)
objectives = [
    ("01", "Python fundamentals", "Apply core syntax and programming concepts.", BLUE, PALE_BLUE),
    ("02", "Control structures", "Use decisions and loops for rules and alerts.", CYAN, PALE_CYAN),
    ("03", "Functions & data", "Organise information using reusable functions and structures.", PURPLE, PALE_PURPLE),
    ("04", "OOP & files", "Practise classes, CSV storage and exception handling.", GREEN, PALE_GREEN),
    ("05", "Data analytics", "Analyse tabular and numerical data with Pandas and NumPy.", AMBER, PALE_AMBER),
    ("06", "NLP concepts", "Explore feedback through transparent text-processing steps.", PURPLE, PALE_PURPLE),
    ("07", "Integrated GUI", "Design a connected desktop application using Tkinter.", BLUE, PALE_BLUE),
]
for i, (num, title, body, col, pale) in enumerate(objectives):
    row, coln = divmod(i, 4)
    x = .65 + coln * 3.15
    y = 1.75 + row * 2.05
    card(s, x, y, 2.9, 1.72, title, body, col, pale, 15, 12)
    circle(s, x + 2.35, y + .1, .38, pale)
    text(s, num, x + 2.35, y + .19, .38, .14, 9, col, True, align=PP_ALIGN.CENTER, margin=0)

# 5. Architecture
s = new_slide("Proposed system architecture", 5)
layers = [
    ("Student & campus data", BLUE, PALE_BLUE),
    ("Input & validation", CYAN, PALE_CYAN),
    ("CSV data storage", GREEN, PALE_GREEN),
    ("Python processing & OOP modules", PURPLE, PALE_PURPLE),
    ("Pandas & NumPy analytics", AMBER, PALE_AMBER),
    ("NLP feedback analysis", CYAN, PALE_CYAN),
    ("Student risk assessment", RED, PALE_RED),
    ("Recommendations & alerts", GREEN, PALE_GREEN),
    ("Tkinter dashboard & reports", BLUE, PALE_BLUE),
]
for i, (lab, col, pale) in enumerate(layers):
    y = 1.58 + i * .58
    node(s, 3.75, y, 5.82, .42, lab, col, pale, 13)
    if i < len(layers) - 1:
        line(s, 6.65, y + .43, 6.65, y + .56, "94A3B8", 1.2)
text(s, "Integrated modules", .78, 2.1, 2.4, .3, 13, NAVY, True)
bullets(s, ["Shared student ID links records.",
            "Validation protects data quality.",
            "Staff review recommendations and alerts."], .78, 2.55, 2.45, 2.0, 12, INK, 10)
text(s, "Proposed architecture; implementation status must be updated by the team.", 3.2, 6.95, 7.0, .18, 9, MUTED, False, align=PP_ALIGN.CENTER, margin=0)

# 6. Technologies
s = new_slide("Technologies used", 6)
tech = [
    ("PY", "Python", "Main programming language.", BLUE, PALE_BLUE),
    ("Tk", "Tkinter + ttk", "Desktop graphical interface.", CYAN, PALE_CYAN),
    ("Pd", "Pandas", "Tabular data processing.", PURPLE, PALE_PURPLE),
    ("Nu", "NumPy", "Numerical calculations.", GREEN, PALE_GREEN),
    ("Mx", "Matplotlib", "Data visualisation.", AMBER, PALE_AMBER),
    ("CSV", "CSV files", "Simple persistent file storage.", BLUE, PALE_BLUE),
    ("NLP", "NLP techniques", "Text cleaning, tokens, keywords and sentiment.", PURPLE, PALE_PURPLE),
]
for i, (abbr, title, body, col, pale) in enumerate(tech):
    row, coln = divmod(i, 4)
    x, y = .68 + coln * 3.12, 1.75 + row * 2.05
    rect(s, x, y, 2.85, 1.68, WHITE, radius=True)
    circle(s, x + .2, y + .22, .58, pale)
    text(s, abbr, x + .2, y + .4, .58, .18, 11, col, True, align=PP_ALIGN.CENTER, margin=0)
    text(s, title, x + .92, y + .27, 1.7, .3, 15, NAVY, True, margin=0)
    text(s, body, x + .22, y + .95, 2.4, .48, 12, INK, margin=0)
text(s, "Technology labels are editable text; no unofficial logos are used.", .75, 6.18, 9.5, .25, 10, MUTED, False)

# 7. Student management
s = new_slide("Student information management", 7)
bullets(s, ["Add, view, search, update and delete records.",
            "Search by student ID or name.",
            "Validate inputs and prevent duplicate IDs.",
            "Confirm before deleting a record."], .72, 1.77, 4.3, 2.65, 15, INK, 11)
rect(s, 5.2, 1.72, 7.45, 3.9, WHITE, radius=True)
text(s, "ILLUSTRATIVE INTERFACE DESIGN • NOT A REAL APPLICATION SCREENSHOT", 5.48, 1.95, 6.9, .23, 9, AMBER, True, margin=0)
headers = ["Student ID", "Name", "Course", "Sem.", "Attendance", "Marks"]
xs = [5.45, 6.58, 8.18, 9.4, 10.15, 11.33]
widths = [1.05, 1.5, 1.15, .62, 1.08, .85]
rect(s, 5.43, 2.42, 6.9, .45, NAVY, radius=True)
for x, w, h in zip(xs, widths, headers):
    text(s, h, x, 2.55, w, .17, 9, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
rows = [
    ["[ID]", "[Name]", "BCA", "3", "—", "—"],
    ["[ID]", "[Name]", "BCA", "3", "—", "—"],
    ["[ID]", "[Name]", "BCA", "3", "—", "—"],
]
for r, row in enumerate(rows):
    y = 3.05 + r * .67
    rect(s, 5.44, y, 6.88, .53, LIGHT, radius=True)
    for x, w, val in zip(xs, widths, row):
        text(s, val, x, y + .16, w, .18, 10, MUTED if "[" in val or val == "—" else INK,
             False, align=PP_ALIGN.CENTER, margin=0)
pill(s, "Student ID", .75, 4.95, 1.35, BLUE, PALE_BLUE)
pill(s, "Name", 2.25, 4.95, .95, CYAN, PALE_CYAN)
pill(s, "Course", 3.35, 4.95, 1.05, PURPLE, PALE_PURPLE)
text(s, "Additional fields: Semester, Email, Attendance, Marks", .78, 5.52, 4.2, .5, 12, MUTED)

# 8. Attendance
s = new_slide("Attendance monitoring", 8)
text(s, "Attendance % = (Attended Classes / Total Classes) × 100", .78, 1.62, 7.0, .45, 20, NAVY, True)
text(s, "Hypothetical example", .8, 2.35, 2.3, .25, 11, MUTED, True)
text(s, "28 / 40 × 100 = 70%", .8, 2.72, 4.0, .5, 25, BLUE, True)
pill(s, "BELOW 75% THRESHOLD", .82, 3.42, 2.5, RED, PALE_RED)
bullets(s, ["Record total and attended classes.",
            "Calculate percentage and display summaries.",
            "Flag records below the 75% threshold.",
            "Generate an attendance alert for staff review."], .8, 4.0, 5.0, 1.8, 14, INK, 8)
rect(s, 7.1, 1.85, 4.95, 3.95, WHITE, radius=True)
text(s, "ATTENDANCE GAUGE • HYPOTHETICAL", 7.42, 2.15, 4.3, .22, 10, MUTED, True, align=PP_ALIGN.CENTER)
circle(s, 8.55, 2.65, 2.05, PALE_BLUE)
circle(s, 8.82, 2.92, 1.51, BLUE)
circle(s, 9.05, 3.15, 1.05, WHITE)
text(s, "70%", 9.03, 3.48, 1.1, .32, 21, NAVY, True, align=PP_ALIGN.CENTER, margin=0)
rect(s, 7.55, 5.0, 4.05, .22, "E2E8F0", radius=True)
rect(s, 7.55, 5.0, 2.84, .22, BLUE, radius=True)
text(s, "40 total classes  •  28 attended", 7.55, 5.33, 4.05, .25, 11, MUTED, False, align=PP_ALIGN.CENTER)

# 9. Academic
s = new_slide("Academic performance analysis", 9)
text(s, "Illustrative marks only • project-defined categories", .72, 1.55, 8.5, .28, 12, MUTED, True)
categories = [("Excellent", "85–100", GREEN, PALE_GREEN), ("Good", "70–84.99", BLUE, PALE_BLUE),
              ("Average", "50–69.99", AMBER, PALE_AMBER), ("Needs Improvement", "Below 50", RED, PALE_RED)]
for i, (name, rng, col, pale) in enumerate(categories):
    x = .72 + i * 3.1
    rect(s, x, 2.0, 2.8, .82, pale, radius=True)
    text(s, name, x + .14, 2.14, 2.5, .25, 14, col, True, align=PP_ALIGN.CENTER)
    text(s, rng, x + .14, 2.47, 2.5, .18, 11, NAVY, False, align=PP_ALIGN.CENTER)
rect(s, .72, 3.22, 7.7, 2.65, WHITE, radius=True)
text(s, "HYPOTHETICAL SUBJECT MARKS", 1.0, 3.48, 6.9, .2, 10, MUTED, True)
for i, (name, score, col) in enumerate([("Python", 78, BLUE), ("Maths", 64, CYAN), ("English", 88, PURPLE), ("DBMS", 52, AMBER)]):
    y = 3.92 + i * .43
    text(s, name, 1.05, y, 1.3, .2, 11, INK, True, margin=0)
    rect(s, 2.55, y + .01, 4.45, .16, "E2E8F0", radius=True)
    rect(s, 2.55, y + .01, 4.45 * score / 100, .16, col, radius=True)
    text(s, str(score), 7.1, y - .02, .45, .2, 11, NAVY, True, align=PP_ALIGN.RIGHT, margin=0)
card(s, 8.85, 3.22, 3.75, 2.65, "Planned analysis",
     "Store subject marks\nCalculate averages when multiple marks exist\nIdentify weak subjects\nSuggest improvement actions",
     PURPLE, PALE_PURPLE, 16, 13)

# 10. Feedback
s = new_slide("Student feedback management", 10)
rect(s, .75, 1.72, 5.2, 4.3, WHITE, radius=True)
text(s, "PROPOSED FEEDBACK FORM", 1.02, 2.0, 4.5, .25, 11, BLUE, True)
label_value(s, "Student ID", "[Enter ID]", 1.02, 2.48, 4.4)
label_value(s, "Date", "[Select date]", 1.02, 3.23, 4.4)
label_value(s, "Concern category", "Academic / Support / Campus / General", 1.02, 3.98, 4.4)
rect(s, 1.02, 4.88, 4.45, .68, LIGHT, radius=True)
text(s, "Feedback text…", 1.2, 5.08, 3.7, .2, 12, MUTED, False, margin=0)
pill(s, "SAVE FEEDBACK", 1.03, 5.72, 1.65, GREEN, PALE_GREEN)
text(s, "Store feedback with student ID and date; keep searchable history.", 6.55, 1.82, 5.3, .52, 17, NAVY, True)
cats = [("Academic", BLUE, PALE_BLUE), ("Support", PURPLE, PALE_PURPLE),
        ("Campus", CYAN, PALE_CYAN), ("General", AMBER, PALE_AMBER)]
for i, (name, col, pale) in enumerate(cats):
    x = 6.55 + (i % 2) * 2.75
    y = 2.72 + (i // 2) * 1.2
    node(s, x, y, 2.35, .76, name, col, pale, 15)
text(s, "Illustrative workflow; not a screenshot of a working application.", 6.55, 5.45, 5.5, .3, 11, MUTED)

# 11. NLP
s = new_slide("NLP-based feedback analysis", 11)
text(s, "TRANSPARENT, RULE-BASED NLP • NOT A TRAINED MODEL UNLESS IMPLEMENTED", .72, 1.48, 11.8, .25, 10, PURPLE, True)
steps = ["Clean", "Lowercase", "Tokenise", "Find keywords", "Classify sentiment", "Identify concern", "Suggest action"]
for i, step in enumerate(steps):
    x = .62 + i * 1.79
    node(s, x, 2.0, 1.55, .64, step, [BLUE, CYAN, PURPLE, AMBER, GREEN, RED, BLUE][i],
         [PALE_BLUE, PALE_CYAN, PALE_PURPLE, PALE_AMBER, PALE_GREEN, PALE_RED, PALE_BLUE][i], 11)
    if i < len(steps) - 1:
        line(s, x + 1.57, 2.32, x + 1.75, 2.32, MUTED, 1.5)
rect(s, .75, 3.05, 4.55, 2.72, WHITE, radius=True)
text(s, "HYPOTHETICAL FEEDBACK", 1.02, 3.32, 3.9, .22, 10, MUTED, True)
text(s, "“I am struggling to understand Python and need help.”", 1.02, 3.78, 3.95, .76, 17, NAVY, True)
rect(s, 5.65, 3.05, 6.9, 2.72, NAVY, radius=True)
text(s, "ILLUSTRATIVE ANALYSIS", 5.95, 3.32, 4.5, .2, 10, "93C5FD", True)
text(s, "Category: Academic", 5.95, 3.78, 5.5, .27, 15, WHITE, True)
text(s, "Concern: Difficulty understanding a subject", 5.95, 4.25, 5.9, .3, 14, WHITE)
text(s, "Suggested action: Faculty mentoring or academic support", 5.95, 4.78, 5.95, .5, 14, "BAE6FD")
pill(s, "SENTIMENT: NEGATIVE", 5.95, 5.33, 1.95, RED, PALE_RED)
pill(s, "POSITIVE", 8.12, 5.33, 1.05, GREEN, PALE_GREEN)
pill(s, "NEUTRAL", 9.32, 5.33, 1.05, BLUE, PALE_BLUE)

# 12. Risk
s = new_slide("AI-based student risk analysis", 12)
text(s, "Project-defined decision-support indicators — not definitive predictions about a student's future.",
     .72, 1.52, 11.9, .38, 13, MUTED, True)
inputs = [("Attendance %", BLUE, PALE_BLUE), ("Academic performance", PURPLE, PALE_PURPLE),
          ("Feedback sentiment", CYAN, PALE_CYAN), ("Repeated concerns", AMBER, PALE_AMBER)]
for i, (lab, col, pale) in enumerate(inputs):
    y = 2.15 + i * .75
    node(s, .85, y, 2.7, .53, lab, col, pale, 13)
    line(s, 3.62, y + .26, 5.35, 3.45, "94A3B8", 1.2)
rect(s, 5.4, 2.78, 2.4, 1.45, NAVY, radius=True)
text(s, "DOCUMENTED\nCONFIGURABLE\nRULES", 5.63, 3.03, 1.95, .8, 14, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line(s, 7.82, 3.5, 8.72, 3.5, BLUE, 2)
risk = [("Low Risk", GREEN, PALE_GREEN), ("Moderate Risk", AMBER, PALE_AMBER),
        ("High Risk", "F97316", "FFEDD5"), ("Support Required", RED, PALE_RED)]
for i, (lab, col, pale) in enumerate(risk):
    y = 2.0 + i * .88
    node(s, 8.9, y, 3.25, .6, lab, col, pale, 14)
text(s, "Staff review is required before action.", 8.92, 5.77, 3.3, .28, 12, MUTED, True, align=PP_ALIGN.CENTER)

# 13. Recommendations
s = new_slide("Intelligent recommendations and alerts", 13)
mapping = [
    ("Low attendance", "Attendance counselling", BLUE, PALE_BLUE),
    ("Low marks", "Subject-specific improvement plan", PURPLE, PALE_PURPLE),
    ("Difficulty understanding a subject", "Faculty mentoring", CYAN, PALE_CYAN),
    ("Repeated campus complaints", "Review the reported concern", AMBER, PALE_AMBER),
    ("Multiple concerns", "Personalised student support review", RED, PALE_RED),
]
for i, (issue, action, col, pale) in enumerate(mapping):
    y = 1.65 + i * .91
    node(s, .92, y, 4.25, .58, issue, col, pale, 13)
    line(s, 5.22, y + .29, 6.55, y + .29, MUTED, 1.6)
    node(s, 6.68, y, 5.4, .58, action, GREEN, PALE_GREEN, 13)
text(s, "Alerts highlight items that may need attention; staff assess context and decide next steps.",
     .95, 6.35, 11.3, .32, 13, NAVY, True, align=PP_ALIGN.CENTER)

# 14. Analytics
s = new_slide("Data analytics and visualisation", 14)
text(s, "Illustrative dashboard concepts • sample values are hypothetical", .72, 1.5, 10.5, .27, 11, MUTED, True)
for i, (name, val, col, pale) in enumerate([("Avg attendance", "—", BLUE, PALE_BLUE),
                                             ("Avg marks", "—", PURPLE, PALE_PURPLE),
                                             ("Risk review", "—", AMBER, PALE_AMBER)]):
    x = .72 + i * 2.1
    rect(s, x, 1.95, 1.9, .95, WHITE, radius=True)
    text(s, name, x + .13, 2.1, 1.62, .2, 10, MUTED, True, align=PP_ALIGN.CENTER)
    text(s, val, x + .13, 2.4, 1.62, .3, 20, col, True, align=PP_ALIGN.CENTER)
rect(s, .72, 3.15, 5.3, 2.65, WHITE, radius=True)
text(s, "HYPOTHETICAL SENTIMENT MIX", .98, 3.4, 4.7, .2, 10, MUTED, True)
for i, (lab, pct, col) in enumerate([("Positive", 45, GREEN), ("Neutral", 30, BLUE), ("Negative", 25, RED)]):
    y = 3.9 + i * .52
    text(s, lab, 1.0, y, 1.25, .2, 11, INK, True)
    rect(s, 2.35, y + .02, 2.8, .16, "E2E8F0", radius=True)
    rect(s, 2.35, y + .02, 2.8 * pct / 100, .16, col, radius=True)
    text(s, f"{pct}%", 5.2, y, .45, .2, 10, MUTED, True)
card(s, 6.4, 1.95, 2.8, 1.62, "Pandas", "Tabular summaries and grouping.", BLUE, PALE_BLUE, 15, 12)
card(s, 9.55, 1.95, 2.8, 1.62, "NumPy", "Numerical calculations.", PURPLE, PALE_PURPLE, 15, 12)
card(s, 6.4, 3.85, 2.8, 1.62, "Matplotlib", "Charts for patterns and distribution.", CYAN, PALE_CYAN, 15, 12)
card(s, 9.55, 3.85, 2.8, 1.62, "Planned views", "Attendance, marks, risk, sentiment and frequent concerns.", AMBER, PALE_AMBER, 15, 11)
text(s, "Charts are editable shapes; no actual project data is represented.", .75, 6.15, 8.2, .23, 10, MUTED)

# 15. Storage
s = new_slide("Data storage and file management", 15)
rect(s, .78, 2.45, 3.0, 1.5, NAVY, radius=True)
text(s, "PYTHON\nAPPLICATION", 1.12, 2.82, 2.3, .68, 20, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
files = ["students.csv", "attendance.csv", "marks.csv", "feedback.csv", "risk_analysis.csv", "recommendations.csv"]
for i, name in enumerate(files):
    coln, row = i % 3, i // 3
    x, y = 5.0 + coln * 2.53, 1.95 + row * 1.65
    node(s, x, y, 2.25, .68, name, [BLUE, CYAN, PURPLE, AMBER, RED, GREEN][i],
         [PALE_BLUE, PALE_CYAN, PALE_PURPLE, PALE_AMBER, PALE_RED, PALE_GREEN][i], 11)
    line(s, 3.82, 3.2, x, y + .34, "94A3B8", 1)
text(s, "Storage workflow", .82, 4.55, 2.5, .28, 13, NAVY, True)
bullets(s, ["Load saved information and save new or updated records.",
            "Handle missing files and empty feedback.",
            "Validate types and incorrect data; protect existing data."], .82, 4.95, 11.5, 1.35, 14, INK, 8)

# 16. GUI/Login
s = new_slide("Graphical user interface and login", 16)
text(s, "PROPOSED WIREFRAMES • NOT WORKING-APPLICATION SCREENSHOTS", .72, 1.47, 11.8, .24, 10, AMBER, True)
rect(s, .72, 1.88, 3.0, 4.05, NAVY, radius=True)
text(s, "CAMPUS SUPPORT", 1.02, 2.25, 2.35, .28, 14, WHITE, True, align=PP_ALIGN.CENTER)
text(s, "Login concept", 1.03, 2.8, 2.3, .34, 17, "BAE6FD", True, align=PP_ALIGN.CENTER)
for yy, lab in [(3.38, "Username"), (4.15, "Password")]:
    rect(s, 1.02, yy, 2.35, .48, WHITE, radius=True)
    text(s, f"[{lab}]", 1.2, yy + .14, 1.9, .18, 11, MUTED, margin=0)
rect(s, 1.02, 4.95, 2.35, .5, BLUE, radius=True)
text(s, "Sign in", 1.2, 5.1, 1.95, .2, 12, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
text(s, "Demo credentials are not specified here because no implemented app was provided.", .98, 5.65, 2.45, .5, 9, "CBD5E1", align=PP_ALIGN.CENTER)
rect(s, 4.12, 1.88, 8.45, 4.05, WHITE, radius=True)
rect(s, 4.12, 1.88, 1.6, 4.05, NAVY, radius=True)
text(s, "MENU", 4.38, 2.2, 1.05, .22, 10, "93C5FD", True, align=PP_ALIGN.CENTER)
for i, lab in enumerate(["Home", "Students", "Attendance", "Academics", "Feedback", "Reports"]):
    text(s, lab, 4.32, 2.68 + i * .43, 1.22, .22, 10, WHITE, i == 0, align=PP_ALIGN.CENTER)
text(s, "Dashboard", 5.98, 2.15, 5.9, .35, 20, NAVY, True)
for i, (lab, col, pale) in enumerate([("Students", BLUE, PALE_BLUE), ("Attendance", CYAN, PALE_CYAN),
                                       ("Support review", AMBER, PALE_AMBER)]):
    x = 5.98 + i * 2.02
    rect(s, x, 2.75, 1.78, .8, pale, radius=True)
    text(s, lab, x + .08, 2.92, 1.62, .2, 10, col, True, align=PP_ALIGN.CENTER)
    text(s, "Planned", x + .08, 3.2, 1.62, .18, 10, MUTED, align=PP_ALIGN.CENTER)
rect(s, 5.98, 3.9, 5.98, 1.45, LIGHT, radius=True)
text(s, "Proposed screens: Student management • Attendance & marks • Feedback/NLP • Risk • Recommendations • Reports • Logout",
     6.25, 4.23, 5.45, .75, 12, INK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# 17. Testing
s = new_slide("Testing and exception handling", 17)
cases = [
    ("Login", "Valid / invalid login"),
    ("Student records", "Duplicate IDs • incorrect data types"),
    ("Marks", "Invalid marks"),
    ("Attendance", "Zero total classes"),
    ("CSV files", "Missing files • save/load records"),
    ("Feedback + NLP", "Empty text • category / sentiment"),
    ("Risk rules", "Risk classification cases"),
]
for i, (area, check) in enumerate(cases):
    y = 1.63 + i * .62
    rect(s, .8, y, 11.75, .48, WHITE, radius=True)
    circle(s, 1.02, y + .12, .23, [BLUE, CYAN, PURPLE, AMBER, GREEN, RED, NAVY][i])
    text(s, area, 1.43, y + .12, 2.35, .2, 12, NAVY, True, margin=0)
    text(s, check, 4.05, y + .12, 7.95, .2, 12, INK, False, margin=0)
text(s, "Planned test cases — no test results are claimed.", .82, 6.25, 9.0, .28, 12, AMBER, True)
text(s, "Show clear error messages and protect existing data.", .82, 6.58, 9.8, .26, 12, MUTED, False)

# 18. Outcomes / scope
s = new_slide("Expected outcomes and future scope", 18)
rect(s, .75, 1.75, 5.75, 4.65, WHITE, radius=True)
rect(s, 6.83, 1.75, 5.75, 4.65, NAVY, radius=True)
pill(s, "EXPECTED OUTCOMES", 1.05, 2.05, 2.2, GREEN, PALE_GREEN)
bullets(s, ["Centralised student information management.",
            "Faster attendance calculations.",
            "Easier academic monitoring and structured feedback analysis.",
            "Explainable recommendations and clear analytical reports."], 1.05, 2.65, 4.95, 2.85, 14, INK, 12)
pill(s, "FUTURE SCOPE", 7.15, 2.05, 1.75, CYAN, PALE_CYAN)
bullets(s, ["Advanced NLP models.",
            "Web or mobile access and database integration.",
            "Role-based access control and historical trends.",
            "Integration with approved campus information systems."], 7.15, 2.65, 4.9, 2.85, 14, WHITE, 12)
text(s, "Expected outcomes and future enhancements are not claims of completed functionality.", 1.0, 5.9, 11.3, .28, 11, MUTED, True, align=PP_ALIGN.CENTER)

# 19. Team and timeline
s = new_slide("Team contribution and project plan", 19)
roles = [
    ("Student 1 • Python Developer", "Student management, control structures, data structures and functions.", BLUE, PALE_BLUE),
    ("Student 2 • System Developer", "OOP-based attendance and academic performance modules.", CYAN, PALE_CYAN),
    ("Student 3 • Data Analytics", "File handling, exceptions, Pandas, NumPy and analysis.", PURPLE, PALE_PURPLE),
    ("Student 4 • AI & GUI Developer", "NLP, risk rules, recommendations, Tkinter, testing and integration.", AMBER, PALE_AMBER),
]
for i, (title, body, col, pale) in enumerate(roles):
    y = 1.58 + i * .77
    card(s, .7, y, 5.85, .65, title, body, col, pale, 13, 10)
text(s, "Team must replace role descriptions with actual names and contributions.", .78, 4.83, 5.65, .35, 10, MUTED, True)
text(s, "14-WEEK PLAN", 7.0, 1.62, 4.8, .25, 11, BLUE, True)
timeline = [
    ("W1–2", "Problem & team"), ("W3–4", "Proposal"), ("W5–6", "Design"),
    ("W7–10", "Development"), ("W11–12", "Testing"), ("W13", "Integration"),
    ("W14", "Demo & submit")
]
for i, (week, event) in enumerate(timeline):
    y = 2.08 + i * .55
    circle(s, 7.08, y + .05, .2, [BLUE, CYAN, PURPLE, GREEN, AMBER, RED, NAVY][i])
    if i < len(timeline) - 1:
        line(s, 7.18, y + .26, 7.18, y + .6, "CBD5E1", 1.6)
    text(s, week, 7.48, y, 1.25, .22, 11, NAVY, True)
    text(s, event, 8.8, y, 3.2, .22, 11, INK)

# 20. Closing
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(NAVY)
rect(s, 0, 0, 13.333, .14, CYAN)
for x, y, d, c in [(10.45, 1.25, .2, CYAN), (11.75, 2.05, .16, PURPLE), (9.65, 3.0, .18, GREEN),
                   (12.0, 4.0, .2, AMBER), (10.55, 5.25, .16, BLUE)]:
    circle(s, x, y, d, c)
for p1, p2 in [((10.55, 1.35), (11.83, 2.13)), ((11.83, 2.13), (9.74, 3.08)),
               ((9.74, 3.08), (12.1, 4.1)), ((12.1, 4.1), (10.63, 5.32))]:
    line(s, p1[0], p1[1], p2[0], p2[1], "64748B", 1.3)
text(s, "CONCLUSION", .85, .75, 4.5, .3, 11, "93C5FD", True, margin=0)
text(s, "A practical Python learning project", .85, 1.25, 8.4, .72, 29, WHITE, True, margin=0)
bullets(s, ["Combines Python, OOP, file handling, analytics, NLP and a GUI.",
            "Designed to support campus monitoring and staff-reviewed support decisions.",
            "An educational prototype for applying programming to a real-world-inspired problem."],
        .9, 2.4, 8.2, 2.25, 16, WHITE, 14)
rect(s, .88, 5.12, 6.0, .04, CYAN)
text(s, "THANK YOU", .88, 5.48, 4.8, .55, 26, "BAE6FD", True, margin=0)
text(s, "Questions and Answers", .9, 6.1, 4.8, .35, 16, WHITE, False, margin=0)
text(s, "Team: [Add names]     Faculty guide: [Add name]", 7.7, 6.17, 4.6, .3, 11, "CBD5E1", False, align=PP_ALIGN.RIGHT)
footer(s, 20, True)

prs.core_properties.title = "AI-Enabled Smart Campus Monitoring and Intelligent Student Support System"
prs.core_properties.subject = "GM University BCA Semester 3 Python PBL"
prs.core_properties.author = "GM University BCA PBL Team"
prs.core_properties.keywords = "Python, Smart Campus, NLP, Data Analytics, PBL"
prs.save(OUT)
print(f"Created {OUT} with {len(prs.slides)} editable slides")
